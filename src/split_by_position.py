# -*- coding: utf-8 -*-
"""按岗位把母题库拆成独立 Excel（每个岗位一份，含该岗位的全部 A/B/C 类题目）。

岗位归属沿用 bank_loader 的推导结果（章节名 -> 岗位），本脚本不做任何硬编码：
A 类章节属于全部岗位，B/C 类章节按 config.xlsx 的部门分组归属。

两种输出格式（--format）：
  kaoshibao（默认）  套用考试宝批量导入模板 templates/kaoshibaoExcel20221101.xlsx
                     保留模板的第 1 行导入须知、第 2 行蓝色表头、「版本号」sheet 与下拉校验，
                     只替换数据行（从第 3 行起）。
  plain              母题库原始列结构（章节/源题号/题型/题干/选项A-H/正确答案/解析/需复核）

用法
    python src/split_by_position.py                      # 考试宝格式 -> data/by_position/
    python src/split_by_position.py --format plain       # 母题库列结构
    python src/split_by_position.py --outdir 目录名 --template 模板路径

输出
    <outdir>/<岗位名>.xlsx     每个岗位一份
    <outdir>/题量汇总.xlsx      岗位 × 题型/类别的题量统计
"""
import argparse
import sys
from collections import Counter
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bank_loader import BASE, load_bank   # noqa: E402

TYPE_NAME = {0: '单选题', 1: '多选题', 2: '判断题'}
CLASSES = ('A类', 'B类', 'C类')

# 母题库原始列
PLAIN_HEADERS = ['章节', '源题号', '题型', '题干',
                 '选项A', '选项B', '选项C', '选项D', '选项E', '选项F', '选项G', '选项H',
                 '正确答案', '解析', '需复核']
PLAIN_WIDTHS = [22, 8, 9, 62, 26, 26, 26, 26, 26, 26, 26, 26, 10, 30, 14]

# 考试宝模板数据列（第 2 行表头由模板自带，这里只按同样顺序写值）
KB_COLUMNS = 14

TEMPLATE_DEFAULT = Path(BASE) / 'templates' / 'kaoshibaoExcel20221101.xlsx'
THIN = Side(style='thin', color='BBBBBB')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def class_of(chapter):
    for c in CLASSES:
        if chapter.endswith(c):
            return c
    return '未分类'


def kb_row(q, chapters):
    """按考试宝模板的 14 列组织一行。

    判断题：模板示例把选项写成「对 / 错」，正确答案用字母 A / B 指向它们
    （模板说明里允许「正确/错误」，但示例统一用字母，遵循示例最稳）。
    """
    if q['t'] == 2:
        opts = ['对', '错'] + [''] * 6
        ans = 'A' if q['a'] == '对' else 'B'
    else:
        opts = list(q['o'])[:8] + [''] * (8 - len(q['o']))
        ans = q['a']
    return [q['s'], TYPE_NAME[q['t']], *opts[:8], ans, q['n'], chapters[q['ch']], '']


def plain_row(q, chapters, num_map):
    if q['t'] == 2:
        opts = ['对', '错'] + [''] * 6
    else:
        opts = list(q['o'])[:8] + [''] * (8 - len(q['o']))
    return [chapters[q['ch']], num_map.get((chapters[q['ch']], q['s']), ''), TYPE_NAME[q['t']], q['s'],
            *opts[:8], q['a'], q['n'], '；'.join(q['f'])]


def write_plain(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '母题库'
    ws.append(PLAIN_HEADERS)
    for r in rows:
        ws.append(r)
    for c in range(1, len(PLAIN_HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = PatternFill('solid', fgColor='4472C4')
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.border = BORDER
    for i in range(2, ws.max_row + 1):
        for c in range(1, len(PLAIN_HEADERS) + 1):
            cell = ws.cell(row=i, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    from string import ascii_uppercase
    for col, w in zip(ascii_uppercase[:len(PLAIN_HEADERS)], PLAIN_WIDTHS):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    wb.save(path)


def write_kaoshibao(template, path, rows):
    """复制模板、只替换数据行，最大程度保留模板形态（须知行/表头/版本号/数据校验）。"""
    wb = openpyxl.load_workbook(template)
    ws = wb[wb.sheetnames[0]]
    # 取第 3 行作为数据行的样式样本
    sample = [copy(ws.cell(row=3, column=c)._style) for c in range(1, KB_COLUMNS + 1)]
    # 清空模板自带的示例数据
    for r in range(3, max(ws.max_row, 3) + 1):
        for c in range(1, KB_COLUMNS + 1):
            ws.cell(row=r, column=c).value = None
    # 去掉示例行的固定行高，让数据行按内容自适应
    for r in [r for r in list(ws.row_dimensions.keys()) if isinstance(r, int) and r >= 3]:
        del ws.row_dimensions[r]
    # 写入
    for i, row in enumerate(rows):
        r = 3 + i
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c)
            cell.value = v
            cell._style = copy(sample[c - 1])
    # 冻结到表头下方、筛选范围扩到数据尾（模板里这两个范围都锚在示例行上）
    ws.freeze_panes = 'A3'
    if rows:
        ws.auto_filter.ref = 'A2:L%d' % (2 + len(rows))
    else:
        ws.auto_filter.ref = None
    wb.save(path)


def load_num_map():
    src = Path(BASE) / 'data' / 'master_bank.xlsx'
    if not src.exists():
        return {}
    wb = openpyxl.load_workbook(src, read_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    head = {n: i for i, n in enumerate(rows[0])}
    out = {}
    for r in rows[1:]:
        if r[0] and r[2] and r[3]:
            out[(str(r[0]).strip(), str(r[head['题干']]).strip())] = r[head['源题号']]
    return out


def main():
    ap = argparse.ArgumentParser(description='按岗位拆分母题库')
    ap.add_argument('--outdir', default='data/by_position', help='输出目录（相对仓库根）')
    ap.add_argument('--format', choices=('kaoshibao', 'plain'), default='kaoshibao',
                    help='输出格式：kaoshibao=考试宝导入模板（默认），plain=母题库列结构')
    ap.add_argument('--template', default=str(TEMPLATE_DEFAULT), help='考试宝模板路径')
    args = ap.parse_args()

    bank = load_bank()
    chapters, qs, positions = bank['chapters'], bank['q'], bank['positions']
    outdir = Path(BASE) / args.outdir
    outdir.mkdir(parents=True, exist_ok=True)

    if args.format == 'kaoshibao':
        tpl = Path(args.template)
        if not tpl.exists():
            print('模板不存在：%s\n已自动改用 plain 格式' % tpl)
            args.format = 'plain'
    num_map = load_num_map() if args.format == 'plain' else {}

    summary = []
    print('%-12s %8s %8s %8s %8s' % ('岗位', 'A类', 'B类', 'C类', '合计'))
    print('-' * 52)
    for i, p in enumerate(positions):
        mine = [q for q in qs if i in q['p']]
        if args.format == 'kaoshibao':
            rows = [kb_row(q, chapters) for q in mine]
            write_kaoshibao(args.template, outdir / ('%s.xlsx' % p['name']), rows)
        else:
            rows = [plain_row(q, chapters, num_map) for q in mine]
            write_plain(outdir / ('%s.xlsx' % p['name']), rows)
        by_class = Counter(class_of(chapters[q['ch']]) for q in mine)
        by_type = Counter(TYPE_NAME[q['t']] for q in mine)
        summary.append([p['name'], by_class.get('A类', 0), by_class.get('B类', 0), by_class.get('C类', 0),
                        len(mine), by_type.get('单选题', 0), by_type.get('多选题', 0),
                        by_type.get('判断题', 0), by_class.get('未分类', 0)])
        print('%-12s %8d %8d %8d %8d' % (p['name'], by_class.get('A类', 0), by_class.get('B类', 0),
                                         by_class.get('C类', 0), len(mine)))

    # 汇总表（辅助文件，不是导入文件）
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '题量汇总'
    ws.append(['岗位', 'A类', 'B类', 'C类', '合计', '单选题', '多选题', '判断题', '未分类'])
    for r in summary:
        ws.append(r)
    ws.append(['—'] * 9)
    ws.append(['母题库唯一题目', sum(1 for q in qs if q['p']), '', '', len(qs),
               sum(1 for q in qs if q['t'] == 0), sum(1 for q in qs if q['t'] == 1),
               sum(1 for q in qs if q['t'] == 2), ''])
    for c in range(1, 10):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4472C4')
    for col, w in zip('ABCDEFGHI', [18, 8, 8, 8, 8, 10, 10, 10, 8]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    wb.save(outdir / '题量汇总.xlsx')

    print('-' * 52)
    print('唯一题目合计: %d（母题库总题数）' % len(qs))
    print('输出格式:', '考试宝导入模板' if args.format == 'kaoshibao' else '母题库列结构')
    print('输出目录:', outdir)
    print('  · %s' % '、'.join('%s.xlsx' % p['name'] for p in positions))
    print('  · 题量汇总.xlsx')


if __name__ == '__main__':
    main()
