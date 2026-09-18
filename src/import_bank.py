# -*- coding: utf-8 -*-
"""从 data/raw_bank 的原始 Word 题库生成《母题库》初版 Excel。

流程：Word(.doc/.docx) --Word COM 抽文本--> 解析 --> 母题库初版 Excel

用法：
    python src/import_bank.py                    # 生成 data/master_bank.imported.xlsx
    python src/import_bank.py --out 别的.xlsx     # 指定输出
    python src/import_bank.py --refresh          # 强制重抽 Word 文本（默认用缓存）
    python src/import_bank.py --text-only DIR    # 只把 Word 正文导出为 txt（不解析）

约定：
  - 章节归属：部门 = raw_bank 子文件夹名（去掉 4 位年份与「题库」后缀），
    类别 = 文件名里的 A类/B类/C类，章节名 = 部门 + 类别 + '类'。
  - 输出列与母题库一致：章节 / 源题号 / 题型 / 题干 / 选项A-H / 正确答案 / 解析 / 需复核。
  - 判断题的选项列固定写 对 / 错，正确答案写 对 / 错。
  - 本脚本只生成「初版」，**不覆盖** data/master_bank.xlsx。Word 原文是终极真值，
    生成后建议用 `python src/verify_bank.py` 与现有母题库交叉核对，再人工合并。
  - 路径与部门名全部从数据推导，代码不含任何业务信息。
"""
import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE = Path(__file__).resolve().parent.parent
RAW_DIR = BASE / 'data' / 'raw_bank'
TEXT_CACHE = BASE / 'data' / 'tmp' / 'import_text.json'
DEFAULT_OUT = BASE / 'data' / 'master_bank.imported.xlsx'

HEADERS = ['章节', '源题号', '题型', '题干',
           '选项A', '选项B', '选项C', '选项D', '选项E', '选项F', '选项G', '选项H',
           '正确答案', '解析', '需复核']

# ---------- 正则 / 常量 ----------
_KB = (r'(单选题|单项选择题|单项选择|多选题|多项选择题|多项选择|不定项选择题|不定项|'
       r'判断题|填空题|问答题|简答题|论述题|计算题|选择题)')
SECTION_RE = re.compile(r'^\s*[一二三四五六七八九十]+\s*[.、．:：]?\s*[（(]?\s*' + _KB + r'[）)]?\s*[.、．:：。]?\s*$')
BARE_SECTION_RE = re.compile(r'^\s*' + _KB + r'\s*[.、．:：。]?\s*$')

SECTION_TYPE_MAP = {
    '单选题': '单选题', '单项选择题': '单选题', '单项选择': '单选题', '选择题': '单选题',
    '多选题': '多选题', '多项选择题': '多选题', '多项选择': '多选题',
    '判断题': '判断题', '填空题': '填空题',
    '问答题': '问答题', '简答题': '问答题', '论述题': '问答题', '计算题': '问答题',
    '不定项': '不定项选择题', '不定项选择题': '不定项选择题',
}

NUM_RE = re.compile(r'^(\d+)\s*[.、．]')
NUM_LEAD_RE = re.compile(r'^\d{1,3}\s*[.、．]\s*')   # 题干行首题号，输出前去掉
OPT_LINE_RE = re.compile(r'^\s*([A-ZＡ-Ｚ])\s*[.、．:：)）]')
ANS_PICK_RE = re.compile(r'[（\[【(]\s*([A-ZＡ-Ｚ][A-ZＡ-Ｚ、，,\s]*?)\s*[）\]】)]')
ANS_TAIL_RE = re.compile(r'([A-H])\s*[。．]?\s*$')
EMPTY_ANS_TAIL_RE = re.compile(r'[（\[【(]\s*[）\]】)]\s*[。．.]?\s*([A-H]{1,8})\s*[。．]?\s*$')  # 空括号后的裸答案，如「（）。ABD」
BARE_MULTI_ANS_RE = re.compile(r'\s([A-H]{2,8})\s*[。．.]?\s*$')                          # 句中裸多字母答案，如「即   AC   。」
INLINE_OPT_RE = re.compile(r'[。．.，,]?\s*([A-Z])\s*[.、．:：)）]\s*([^，。]{1,80})$')
JUDGE_WORD = r'(正确|错误|对|错|√|×|✓|T|F|是|否)'
ANS_JUDGE_RE = re.compile(r'[（\[【(]\s*(' + JUDGE_WORD + r')\s*[）\]】)]')
ANS_JUDGE_PLAIN_RE = re.compile(r'(?<![A-Za-z0-9\u4e00-\u9fff])(' + JUDGE_WORD + r')\s*$')
ANS_JUDGE_NOTE_RE = re.compile(r'(?<![A-Za-z0-9\u4e00-\u9fff])(' + JUDGE_WORD + r')\s*[（\[【(][^）\]】)]*[）\]】)]\s*$')

FW = str.maketrans('ＡＢＣＤＥＦＧＨＴＦ', 'ABCDEFGHTF')
JUDGE_CANON = {'对': '对', '是': '对', '正确': '对', '√': '对', '✓': '对', 'T': '对',
               '错': '错', '否': '错', '错误': '错', '×': '错', 'F': '错'}

NOTE_PREFIX_RE = re.compile(r'^(解析|（?正确答案）?|答案)\s*[:：]?\s*[为是]?\s*')
TRAIL_EMPTY_PAREN_RE = re.compile(r'\s*[（\[【(]\s*[）\]】)]\s*$')


def norm_fw(s):
    return s.translate(FW)


def opt_like(txt):
    """判断一段文本是否为真实选项内容（去标点后 >=2 字符，或含至少 1 个汉字）"""
    clean = re.sub(r'[\s、，,。．:：()（）【】\[\]"“”\u3000\xa0]+', '', txt)
    if not clean:
        return False
    if len(clean) >= 2:
        return True
    return any('\u4e00' <= c <= '\u9fff' for c in clean)


def has_answer_marker(s):
    s = norm_fw(s)
    return bool(ANS_PICK_RE.search(s) or ANS_JUDGE_RE.search(s))


def has_judge_marker(s):
    s = norm_fw(s)
    return bool(ANS_JUDGE_RE.search(s) or ANS_JUDGE_PLAIN_RE.search(s) or ANS_JUDGE_NOTE_RE.search(s))


def classify_section(title):
    for k in ('单项选择题', '多项选择题', '单选题', '多选题', '不定项选择题', '判断题',
              '填空题', '问答题', '简答题', '论述题', '计算题', '单项选择', '多项选择', '不定项', '选择题'):
        if k in title:
            return SECTION_TYPE_MAP[k]
    return None


def strip_option_noise(txt):
    txt = re.sub(r'[。．、]+$', '', txt.strip()).rstrip()
    return txt.strip('"“”')


def clean_stem(s):
    return re.sub(r'\s{2,}', ' ', s).strip().strip('"“”\u3000 ')


def extract_judge_answer(line):
    """抽取判断题答案，返回 (答案'对'/'错' 或 None, 去掉答案的题干, 括号内补充说明)"""
    m = ANS_JUDGE_RE.search(line)
    if m:
        ans = JUDGE_CANON[norm_fw(m.group(1))]
        stem = TRAIL_EMPTY_PAREN_RE.sub('', line[:m.start()] + line[m.end():]).rstrip()
        return ans, stem, line[m.end():].strip()
    m2 = ANS_JUDGE_PLAIN_RE.search(line)
    if m2:
        ans = JUDGE_CANON[norm_fw(m2.group(1))]
        return ans, TRAIL_EMPTY_PAREN_RE.sub('', line[:m2.start()]).rstrip(), ''
    m3 = ANS_JUDGE_NOTE_RE.search(line)
    if m3:
        ans = JUDGE_CANON[norm_fw(m3.group(1))]
        stem = TRAIL_EMPTY_PAREN_RE.sub('', line[:m3.start()]).rstrip()
        nm = re.search(r'[（\[【(]([^）\]】)]*)[）\]】)]\s*$', line[m3.start():])
        return ans, stem, (nm.group(1).strip() if nm else '')
    return None, line, ''


def parse_option_line(line):
    """按字母标记切分一行中的选项，返回 (leading, [(letter, text), ...])；不是选项行返回 None。
    支持「字母+标点」与「字母+空格」两种分隔。"""
    line = norm_fw(line)
    strict = list(re.finditer(r'(?<![A-Z])([A-Z])\s*[.、．:：)）]', line))
    space = list(re.finditer(r'(?<![A-Z])([A-Z])\s{1,}(?=[^\s，。、])', line))
    matches = None
    if len(strict) >= 2:
        matches = strict
    elif len(strict) == 1:
        # 注意：marker 在行首时无前一位字符，必须单独判断；
        # 不能用 ``prev not in '（(【〔'``——空串是任意字符串的子串，会恒为 False。
        m0 = strict[0]
        if m0.start() == 0 or line[m0.start() - 1] not in '（(【〔':
            matches = strict
    elif len(space) >= 2:
        matches = space
    if not matches:
        return None
    leading = line[:matches[0].start()].strip()
    segs = []
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(line)
        segs.append((m.group(1), line[m.end():end]))
    return leading, segs


def extract_pick_answer(line):
    """抽取选择题答案，返回 (字母列表 或 None, 去掉答案的题干, 内联选项列表)"""
    line = norm_fw(line)
    ms = list(ANS_PICK_RE.finditer(line))
    if ms:
        groups, ok = [], True
        for m in ms:
            l = re.sub(r'[\s、，,．.]+', '', m.group(1)).upper()
            if not l:
                ok = False
                break
            groups.append((l, m))
        if ok and groups:
            singles = [g for g in groups if len(g[0]) == 1]
            pick = singles if singles else [groups[-1]]
            letters = [g[0] for g in pick]
            stem = line
            for g in pick:
                stem = stem[:g[1].start()] + stem[g[1].end():]
            r = parse_option_line(stem)
            if r is not None:
                leading, segs = r
                if (len(segs) >= 2 and opt_like(segs[0][1]) and segs[0][0] == 'A' and leading.strip()):
                    return letters, leading.strip('。．、　 ').rstrip('。．、'), segs
            tail = INLINE_OPT_RE.search(stem)
            if tail:
                return letters, stem[:tail.start()].rstrip('。．'), [(tail.group(1), tail.group(2).strip())]
            return letters, stem, []
    m0 = EMPTY_ANS_TAIL_RE.search(line)          # 先处理「（）ABD」式裸答案
    if m0:
        return list(m0.group(1)), line[:m0.start()].rstrip(), []
    m1 = BARE_MULTI_ANS_RE.search(line)          # 句中裸多字母答案（≥2 个字母，避免误伤单个字母结尾）
    if m1:
        return list(m1.group(1)), line[:m1.start()].rstrip(), []
    m2 = ANS_TAIL_RE.search(line)
    if m2:
        return [m2.group(1)], line[:m2.start()].rstrip(), []
    return None, line, []


def parse_file(raw):
    """把整篇文本切成 (区段题型, [题目块])；题目块保留原始行，稍后 refine。"""
    raw = raw.replace('\u000b', '\n').replace('\r\n', '\n').replace('\r', '\n')
    lines = raw.split('\n')

    sections, cur_type, cur_questions, cur_q, cur_has_ans = [], None, [], None, False

    def marker(s, t):
        return has_judge_marker(s) if t == '判断题' else has_answer_marker(s)

    def flush_q():
        nonlocal cur_q, cur_has_ans
        if cur_q is not None:
            cur_questions.append(cur_q)
            cur_q, cur_has_ans = None, False

    def start_q(src_no, line):
        nonlocal cur_q, cur_has_ans
        flush_q()
        cur_q = {'src_no': src_no, 'raw_lines': [line]}
        cur_has_ans = marker(line, cur_type)

    for ln in lines:
        s = ln.strip()
        if not s:
            continue
        if SECTION_RE.match(s) or BARE_SECTION_RE.match(s):
            t = classify_section(s)
            if t is not None:
                flush_q()
                if cur_questions:
                    sections.append((cur_type, cur_questions))
                cur_type, cur_questions = t, []
                continue
        mn = NUM_RE.match(s)
        if mn:
            start_q(int(mn.group(1)), s)
            continue
        if cur_q is None:
            if marker(s, cur_type):
                start_q(None, s)
            continue
        if cur_has_ans and marker(s, cur_type) and not OPT_LINE_RE.match(s):
            start_q(None, s)
            continue
        cur_q['raw_lines'].append(s)
        if marker(s, cur_type):
            cur_has_ans = True
    flush_q()
    if cur_questions:
        sections.append((cur_type, cur_questions))
    return sections


def refine_question(q, sec_type):
    """把原始行解析成 题干/选项/答案/解析/flags。"""
    lines = q['raw_lines']
    stem_parts, options, answer, answer_src, note, flags = [], [], None, '', '', []

    def next_letter():
        return chr(ord('A') + len(options)) if len(options) < 8 else None

    if sec_type == '判断题':
        for i, ln in enumerate(lines):
            s = ln.strip()
            if not s:
                continue
            if answer is None:
                ans, stem, n = extract_judge_answer(s)
                if ans is not None:
                    answer = ans
                    answer_src = '题干' if i == 0 else '独立行'
                    if stem.strip():
                        stem_parts.append(stem.strip())
                    if n:
                        note = n
                    continue
            if s.startswith(('解析', '（正确答案', '正确答案')):
                note = note or NOTE_PREFIX_RE.sub('', s).strip('（）')
                continue
            stem_parts.append(s)
        if answer is None:
            flags.append('未提取到判断题答案')
    else:
        first = lines[0].strip() if lines else ''
        ans_groups, stem0, tail_opts = extract_pick_answer(first)
        if ans_groups:
            answer = ''.join(dict.fromkeys(''.join(ans_groups)))
            answer_src = '题干'
            if stem0.strip():
                stem_parts.append(stem0.strip())
            for letter, txt in tail_opts:
                options.append((letter, strip_option_noise(txt)))
        else:
            stem_parts.append(first)

        for idx in range(1, len(lines)):
            s = lines[idx].strip()
            if not s:
                continue
            if s.startswith(('解析', '（正确答案', '正确答案', '答案')):
                note = note or NOTE_PREFIX_RE.sub('', s).strip('（）')
                continue
            r = parse_option_line(s)
            if r is not None and (len(r[1]) >= 2 or any(opt_like(t) for _, t in r[1])):
                leading, segs = r
                if leading:
                    letter = next_letter()
                    if letter:
                        options.append((letter, strip_option_noise(leading)))
                for letter, txt in segs:
                    options.append((letter, strip_option_noise(txt)))
                continue
            if answer is None:
                ans_groups, stem2, tail_opts = extract_pick_answer(s)
                if ans_groups:
                    answer = ''.join(dict.fromkeys(''.join(ans_groups)))
                    answer_src = '题干(后续行)'
                    if stem2.strip():
                        stem_parts.append(stem2.strip())
                    for letter, txt in tail_opts:
                        options.append((letter, strip_option_noise(txt)))
                    continue
            # 跨行续写识别：本行不是选项行，且下一行以「尚未占用的选项字母」开头
            # ⇒ 本行是上一个选项的排版换行，应并入上一个选项而不是新开一个。
            used = {l for l, _ in options}
            nxt = norm_fw(lines[idx + 1].strip()) if idx + 1 < len(lines) else ''
            nxt_opt = OPT_LINE_RE.match(nxt)
            if (options and len(s) <= 20 and not re.search(r'[。．；;）)]$', s)
                    and nxt_opt and norm_fw(nxt_opt.group(1)) not in used):
                options[-1] = (options[-1][0], options[-1][1] + s)
                continue
            letter = next_letter()
            if letter:
                options.append((letter, strip_option_noise(s)))
            else:
                flags.append('选项超过8个，无法续接')

    stem = NUM_LEAD_RE.sub('', clean_stem(' '.join(p for p in stem_parts if p)))

    # 选项字母超 H（如 J）时按出现顺序重标为 A-H，并同步答案（考试宝仅支持 A-H）
    if options and any(l not in 'ABCDEFGH' for l, _ in options):
        mapping = {l: chr(ord('A') + i) for i, (l, _) in enumerate(options)}
        options = [(mapping[l], t) for l, t in options]
        if answer and answer_src:
            answer = ''.join(mapping.get(ch, ch) for ch in answer)
        flags.append('选项字母超H，已按出现顺序重标为A-H并同步答案')

    q.update(stem=stem, answer=answer, answer_src=answer_src,
             options=options, note=note, flags=flags, sec_type=sec_type)

    # 源文件标为单选但答案多字母时，改判为多选
    if sec_type == '单选题' and answer and len(answer) > 1:
        q['sec_type'] = '多选题'
        flags.append('源文件标为单选题但答案多选，已按实际类型改为多选题')

    q['valid'] = answer is not None
    if answer is None:
        flags.append('未提取到答案')
    else:
        if q['sec_type'] == '单选题' and len(answer) != 1:
            flags.append(f'单选题答案非单字母: {answer}')
        if q['sec_type'] == '多选题' and len(answer) > len({l for l, _ in options}):
            flags.append(f'多选题答案数超过选项数: {answer}')
    if options and answer and q['sec_type'] in ('单选题', '多选题', '不定项选择题'):
        mg = {l for l, _ in options}
        for al in set(answer):
            if al not in mg:
                flags.append(f'答案字母{al}不在选项A-H中')
    if q['sec_type'] in ('单选题', '多选题', '不定项选择题') and len(options) < 2:
        flags.append('选项少于2个，疑似切题/切选项异常')
    if not stem:
        flags.append('题干为空')
    return q


def infer_type(q):
    """无区段标题时按答案特征推断题型。"""
    for ln in q['raw_lines']:
        if ANS_JUDGE_RE.search(ln) or ANS_JUDGE_PLAIN_RE.search(ln) or ANS_JUDGE_NOTE_RE.search(ln):
            return '判断题'
    for ln in q['raw_lines']:
        m = ANS_PICK_RE.search(ln)
        if m:
            letters = re.sub(r'[\s、，,．.]+', '', m.group(1)).upper()
            return '多选题' if len(letters) > 1 else '单选题'
    return '单选题'


# ---------- Word 抽取 ----------
def extract_docs(refresh):
    """把 data/raw_bank 下全部 Word（.doc/.docx）正文抽成 {相对路径: 文本}，带缓存。"""
    if TEXT_CACHE.exists() and not refresh:
        return json.loads(TEXT_CACHE.read_text(encoding='utf-8')), 'cache'
    try:
        import win32com.client as win32
    except ImportError:
        sys.exit('未安装 pywin32，无法抽取 Word 正文。请先 pip install pywin32，或用 --text-only 提供文本。')
    files = sorted(p for p in list(RAW_DIR.glob('*/*.doc')) + list(RAW_DIR.glob('*/*.docx'))
                   if not p.name.startswith('~$'))
    app = win32.gencache.EnsureDispatch('Word.Application')
    app.Visible = False
    app.DisplayAlerts = 0
    dump = {}
    try:
        for f in files:
            rel = f.relative_to(RAW_DIR).as_posix()
            doc = app.Documents.Open(str(f), ReadOnly=True, AddToRecentFiles=False)
            text = doc.Content.Text
            doc.Close(False)
            dump[rel] = text.replace('\x07', '\n').replace('\x0b', '\n').replace('\x0c', '\n')
    finally:
        app.Quit()
    TEXT_CACHE.parent.mkdir(parents=True, exist_ok=True)
    TEXT_CACHE.write_text(json.dumps(dump, ensure_ascii=False), encoding='utf-8')
    return dump, '%d 个文件' % len(dump)


def derive_chapter(rel_path):
    """由 raw_bank 相对路径推导章节名：部门=文件夹名，类别=文件名里的 A/B/C 类。"""
    p = Path(rel_path)
    dept = re.sub(r'^\d{4}年?', '', p.parent.name)
    dept = re.sub(r'题库$', '', dept).strip()
    m = re.search(r'([ABC])类', p.name)
    if not dept or not m:
        return None
    return f'{dept}{m.group(1)}类'


# ---------- 写 Excel ----------
def write_excel(out_path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '母题库'
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)

    thin = Side(style='thin', color='BBBBBB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill('solid', fgColor='4472C4')
    hdr_font = Font(bold=True, color='FFFFFF')
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = hdr_fill, hdr_font
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for i in range(2, ws.max_row + 1):
        for c in range(1, len(HEADERS) + 1):
            ws.cell(row=i, column=c).border = border
    for col, w in zip('ABCDEFGHIJKLMNO', [14, 7, 9, 46, 20, 20, 20, 20, 20, 20, 20, 20, 8, 28, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    wb.save(out_path)


def to_row(chapter, q):
    """结构化记录 -> Excel 行（判断题选项固定为 对/错）。"""
    sec = q['sec_type']
    if sec == '判断题':
        options = ['对', '错'] + [''] * 6
        ans = q['answer'] or ''
    else:
        options = [t for _, t in q['options']][:8]
        options += [''] * (8 - len(options))
        ans = q['answer'] or ''
    return [chapter, q.get('src_no'), sec, q['stem'], *options,
            ans, q.get('note') or '', '；'.join(q.get('flags') or [])]


def main():
    ap = argparse.ArgumentParser(description='从 raw_bank 的 Word 题库生成母题库初版 Excel')
    ap.add_argument('--out', default=str(DEFAULT_OUT), help='输出 xlsx（默认 data/master_bank.imported.xlsx）')
    ap.add_argument('--refresh', action='store_true', help='强制重抽 Word 正文（忽略缓存）')
    ap.add_argument('--text-only', metavar='DIR', default=None, help='只把 Word 正文导出到该目录（不解析）')
    args = ap.parse_args()

    if not RAW_DIR.exists():
        sys.exit(f'未找到 {RAW_DIR}')

    if args.text_only:
        dump, info = extract_docs(args.refresh)
        outdir = Path(args.text_only)
        outdir.mkdir(parents=True, exist_ok=True)
        for rel, text in dump.items():
            (outdir / (rel.replace('/', '__') + '.txt')).write_text(text, encoding='utf-8')
        print(f'导出 {info} -> {outdir}')
        return

    dump, info = extract_docs(args.refresh)
    print(f'Word 正文：{info}')

    rows, skipped, per_ch = [], [], Counter()
    for rel in sorted(dump):
        chapter = derive_chapter(rel)
        if chapter is None:
            skipped.append(rel)
            continue
        sections = parse_file(dump[rel])
        for sec_type, qs in sections:
            for q in qs:
                t = sec_type if sec_type is not None else None
                if t is None:
                    t = infer_type(q)
                refine_question(q, t)
                if not q.get('stem'):
                    continue
                rows.append(to_row(chapter, q))
                per_ch[chapter] += 1

    if skipped:
        print('跳过（无法推导章节）:', skipped)

    out = Path(args.out)
    write_excel(out, rows)
    print(f'生成 -> {out}（{len(rows)} 题 / {len(per_ch)} 章节）')
    types = Counter(r[2] for r in rows)
    print('题型分布:', dict(types))
    flagged = sum(1 for r in rows if r[14])
    print(f'需复核: {flagged} 题')
    print('\n各章节题量:')
    for ch in sorted(per_ch):
        print(f'  {ch:<20} {per_ch[ch]}')


if __name__ == '__main__':
    main()
