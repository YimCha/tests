# -*- coding: utf-8 -*-
"""一次性：把期货/银行等中文 Word 题库（docx）整理为 Excel。

用法：
    python src/parse_docx_bank.py "其他/2026年运营风采竞技大赛理论.docx"
    python src/parse_docx_bank.py                     # 自动扫描 ../other 下第一个 docx

输入 docx 的题目格式约定（与模板一致）：
    一、单项选择题 / 二、多项选择题 / 三、判断题   （区段标题）
    1. 题干（  ）                               （编号 + 题干，可能带空括号占位）
    A. 选项一  B. 选项二  C. 选项三  D. 选项四    （选项可同行或分行）
    答案：B                                     （每题必有且仅有一行）
判断答案可为 对/错/正确/错误，括号内文字（如 错（需重点关注））作为解析。

输出：与源 docx 同目录生成 <文件名>_整理.xlsx，列结构与母题库对齐：章节/源题号/题型/
题干/选项A-H/正确答案/解析/需复核。
"""
import argparse
import re
import sys
from pathlib import Path

import docx
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# 同一目录下的 other 库，脚本自身定位，不写死绝对路径
SRC_DIR = Path(__file__).resolve().parent.parent / 'other'

SEC_HDR_RE = re.compile(
    r'^[一二三四五六七八九十]+、\s*(.*?)\s*$')
SEC_TYPE_RE = re.compile(r'(单项选择题|多项选择题|判断题|不定项选择题|单项选择题|填空题|简答题|问答题|论述题)')
NUM_RE = re.compile(r'^(\d+)\s*[.、．]\s*')
ANS_LINE_RE = re.compile(r'^答案\s*[:：]?\s*(.*)$')
OPT_SEG_RE = re.compile(r'(?<![A-Za-z])([A-H])\s*[.、．:：)）]')
TRAIL_PAREN_RE = re.compile(r'[（(]\s*[）)]\s*$')   # 题干末尾空括号占位

JUDGE_CANON = {'正确': '对', '对': '对', '是': '对',
               '错误': '错', '错': '错', '否': '错'}
JUDGE_PAT = re.compile(r'^(正确|对|是|错误|错|否)\s*(?:[（(](.*?)[）)])?\s*$')

HEADERS = ['章节', '源题号', '题型', '题干',
           '选项A', '选项B', '选项C', '选项D', '选项E', '选项F', '选项G', '选项H',
           '正确答案', '解析', '需复核']
COL_WIDTHS = [24, 8, 9, 46, 20, 20, 20, 20, 20, 20, 20, 20, 10, 28, 12]

FW = str.maketrans('ＡＢＣＤＥＦＧＨ', 'ABCDEFGH')


def norm_fw(s):
    return s.translate(FW)


def classify_sec(header):
    m = SEC_TYPE_RE.search(header)
    if not m:
        return None
    t = m.group(1)
    if '单项' in t:
        return '单选题'
    if '多项' in t:
        return '多选题'
    if '判断' in t:
        return '判断题'
    return None


def split_lines(doc):
    return [p.text.strip() for p in doc.paragraphs]


def group_questions(lines):
    """按区段标题分组，再按编号把题目切成块。返回 [(题型, [(题号, [本行…])])]"""
    sections = []
    cur_type = None
    cur_qs = []
    for s in lines:
        if not s:
            continue
        m = SEC_HDR_RE.match(s)
        if m and SEC_TYPE_RE.search(s):
            t = classify_sec(s)
            if t is not None:
                if cur_qs:
                    sections.append((cur_type, cur_qs))
                cur_type = t
                cur_qs = []
                continue
        if cur_type is None:
            continue
        nm = NUM_RE.match(s)
        if nm:
            cur_qs.append((int(nm.group(1)), [s]))
        elif cur_qs:
            cur_qs[-1][1].append(s)
    if cur_qs:
        sections.append((cur_type, cur_qs))
    return sections


def split_options(line):
    """按 A-H 字母前缀把一行切成选项段，返回 [(letter, text), ...]。
    首段无字母前缀时不当作选项（判为题干续行）。"""
    line = norm_fw(line)
    ms = list(OPT_SEG_RE.finditer(line))
    if not ms or ms[0].start() != 0:
        return None
    segs = []
    for i, m in enumerate(ms):
        letter = m.group(1)
        end = ms[i + 1].start() if i + 1 < len(ms) else len(line)
        txt = line[m.end():end].strip().rstrip('。.、')
        segs.append((letter, txt))
    return segs


def clean_stem(txt):
    txt = re.sub(r'\s+', ' ', txt).strip('\u00a0 ')
    txt = norm_fw(txt)
    txt = TRAIL_PAREN_RE.sub('', txt).strip()
    return txt


def parse_judge_answer(raw):
    m = JUDGE_PAT.match(raw)
    if not m:
        return None, raw, ''
    ans = JUDGE_CANON[m.group(1)]
    return ans, '', (m.group(2) or '')


def parse_block(sec_type, num, lines):
    stem_lines = [lines[0]]
    options = []
    answer = None
    note = ''
    flags = []
    answer_line_seen = False

    for s in lines[1:]:
        if answer_line_seen:
            stem_lines.append(s)
            continue
        am = ANS_LINE_RE.match(s)
        if am:
            answer_line_seen = True
            if sec_type == '判断题':
                ans, _, exp = parse_judge_answer(am.group(1).strip())
                if ans is not None:
                    answer = ans
                else:
                    flags.append('判断答案无法解析: ' + am.group(1))
                note = exp
            else:
                answer = ''.join(dict.fromkeys(norm_fw(am.group(1)).upper()))
            continue
        segs = split_options(s)
        if segs:
            options.extend(segs)
        else:
            stem_lines.append(s)

    stem = clean_stem(' '.join(stem_lines))

    if sec_type == '判断题':
        option_texts = ['对', '错']
    else:
        option_texts = []
        opt_map = {}
        for letter, txt in options:
            opt_map.setdefault(letter, txt)
        for letter in 'ABCDEFGH':
            option_texts.append(opt_map.get(letter, ''))
    option_texts = (option_texts + [''] * 8)[:8]

    # 校验
    if sec_type in ('单选题', '多选题'):
        ans = answer or ''
        opt_letters = {l for l, _ in options}
        for ch in ans:
            if ch not in opt_letters:
                flags.append(f'答案字母{ch}不在选项A-H中')
        if sec_type == '单选题' and len(ans) > 1:
            flags.append('单选题答案多选: ' + ans)
    if not answer:
        flags.append('未提取到答案')

    row = [stem, *option_texts[:8], answer or '', note, '；'.join(flags)]
    return row


def build_rows(sections):
    rows = []
    for sec_type, qs in sections:
        for num, lines in qs:
            body = parse_block(sec_type, num, lines)
            rows.append([num, sec_type, *body])
    return rows


def write_excel(out_path, chapter, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '题库'
    ws.append(HEADERS)
    for r in rows:
        ws.append([chapter, *r])

    thin = Side(style='thin', color='BBBBBB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill('solid', fgColor='4472C4')
    hdr_font = Font(bold=True, color='FFFFFF')
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for i in range(2, ws.max_row + 1):
        for c in range(1, len(HEADERS) + 1):
            ws.cell(row=i, column=c).border = border
    from string import ascii_uppercase
    for col, w in zip(ascii_uppercase[:len(HEADERS)], COL_WIDTHS):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    wb.save(out_path)
    return len(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('input', nargs='?', default=None,
                    help='docx 路径；缺省自动取 ../other 下第一个 docx')
    args = ap.parse_args()

    if args.input:
        src = Path(args.input)
    else:
        candidates = [p for p in SRC_DIR.glob('*.docx') if not p.name.startswith('~$')]
        if not candidates:
            sys.exit('未找到 docx，请传入文件路径')
        src = candidates[0]

    doc = docx.Document(str(src))
    lines = split_lines(doc)
    sections = group_questions(lines)
    if not sections:
        sys.exit('未解析到任何题目区段，请检查 docx 结构')

    rows = build_rows(sections)
    out = src.with_name(src.stem + '_整理.xlsx')
    total = write_excel(out, src.stem, rows)
    print('区段:', [(t, len(q)) for t, q in sections])
    print('总题数:', total)
    print('输出 ->', out)


if __name__ == '__main__':
    main()