"""按考试方案比例生成岗位试卷导入表（考试模板格式 xlsx）。

从 data/config.xlsx 读岗位方案（分值比例），从 data/master_bank.xlsx 读题库，
每个岗位一张卷：题型题量默认 单选42/多选43/判断15（--types 可改），每题 1 分。
抽题单元（A 合集 / B、C 部门组）题量 = 方案分值 ÷ 80 × 总题量，最大余数法取整；
单元池不足时该单元按实有题数出卷，缺口并入 A 合集（A 合集为全库 A 类题，容量最大）。
单元池互斥（A/B/C 类别×部门组天然不重叠），单元池总量不足时该单元按实有题数出卷、
缺口并入 A 合集；单元内某题型池不足时，由「题型×单元」矩阵在容量约束内把该题型
缺口调剂到其他单元，单元题量与题型题量均精确成立（与 _app_script.js 的 genPaper
同源比例口径，但矩阵求解带容量上限，题型分布不漂移）。
输出列：题型 / 题目 / 选项1~8 / 正确答案 / 答案解析 / 分值；
判断题选项固定「A、正确」「B、错误」，正确答案写 A/B。

用法:
  python src/export_papers.py                          # 全部岗位
  python src/export_papers.py --positions <岗位名 ...>  # 指定岗位
  python src/export_papers.py --types 40 25 15 --seed 2026
输出: data/papers/{岗位名}试卷.xlsx
"""
import argparse
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from bank_loader import load_positions, load_rows

BASE = Path(__file__).resolve().parent.parent
OUT_DIR = BASE / 'data' / 'papers'
TYPE_NAME = {0: '单选题', 1: '多选题', 2: '判断题'}
LETTERS = 'ABCDEFGH'
HEADERS = ['题型', '题目'] + [f'选项{i}' for i in range(1, 9)] + ['正确答案', '答案解析', '分值']


def split_by_weight(total, weights):
    """total 按 weights 比例拆整数（最大余数法），合计恒等于 total。"""
    raw = [total * w / sum(weights) for w in weights]
    res = [int(x) for x in raw]
    order = sorted(range(len(raw)), key=lambda i: raw[i] - res[i], reverse=True)
    for k in range(total - sum(res)):
        res[order[k % len(order)]] += 1
    return res


def _rotate(m, row_left, col_left, caps):
    """贪心卡死时的三点轮换：t2 行在 c 列尚有池容量但行配额已满，
    让出它在 c2 列的一题给仍有行配额的 t3 行，自己补进 c 列。"""
    for c in range(len(col_left)):
        if col_left[c] <= 0:
            continue
        for t2 in range(len(m)):
            if row_left[t2] != 0 or m[t2][c] >= caps[t2][c]:
                continue
            for c2 in range(len(col_left)):
                if m[t2][c2] <= 0:
                    continue
                for t3 in range(len(m)):
                    if row_left[t3] > 0 and m[t3][c2] < caps[t3][c2]:
                        m[t2][c] += 1
                        m[t2][c2] -= 1
                        m[t3][c2] += 1
                        col_left[c] -= 1
                        row_left[t3] -= 1
                        return True
    return False


def gen_matrix(types, targets, caps):
    """[题型×抽题单元] 目标矩阵：行合计=types、列合计=targets，
    每格不超过 caps[t][c]（该单元该题型池大小），比例口径与工具 genMatrix 一致。"""
    total = sum(types)
    m = [[min(int(types[t] * targets[c] / total), caps[t][c])
          for c in range(len(targets))] for t in range(len(types))]
    cells = []
    for t, n in enumerate(types):
        for c, g in enumerate(targets):
            e = n * g / total
            cells.append((t, c, e - int(e)))
    row_left = [types[t] - sum(m[t]) for t in range(len(types))]
    col_left = [targets[c] - sum(m[t][c] for t in range(len(types))) for c in range(len(targets))]
    cells.sort(key=lambda x: -x[2])
    while sum(row_left) > 0:
        progressed = False
        for t, c, _ in cells:
            if row_left[t] > 0 and col_left[c] > 0 and m[t][c] < caps[t][c]:
                m[t][c] += 1
                row_left[t] -= 1
                col_left[c] -= 1
                progressed = True
        if not progressed and not _rotate(m, row_left, col_left, caps):
            raise ValueError('题型×单元矩阵在池容量内无解，请检查题库规模或题型口径')
    return m


def build_units(pos, chapters):
    """抽题单元：A 合集 + B 部门组 + C 部门组，返回 [(标签, 池索引列表, 方案分值)]。"""
    units = [('A合集', [i for i, ch in enumerate(chapters) if ch.endswith('A类')], pos['ratio'][0])]
    for cls, groups in (('B', pos['b']), ('C', pos['c'])):
        for gi, (depts, pts) in enumerate(groups):
            pool = [i for i, ch in enumerate(chapters)
                    if ch.endswith(cls + '类') and ch[:-2] in depts]
            units.append((f'{cls}{gi + 1}:{"+".join(depts)}', pool, pts))
    return units


def gen_paper(pos, rows, types, rng):
    """一个岗位一张卷，返回 (题目索引列表, 各单元题量)。
    同卷题干唯一：题库存在同题干多版本题（选项/答案不同），同一张卷每题干只出其一，
    单元容量与矩阵上限均按去重后的题干数计。"""
    chapters = [ch for ch, _ in rows]
    units = build_units(pos, chapters)
    stem_pools = []  # 每单元按题型分组的 {题干: [题目索引, ...]}
    seen_stems = set()
    for _, pool, _ in units:
        by_t = [{} for _ in types]
        for i in pool:
            q = rows[i][1]
            if q['s'] not in seen_stems:  # 跨单元同题干只归最先遇到的单元，保证全卷题干唯一
                by_t[q['t']].setdefault(q['s'], []).append(i)
        for d in by_t:
            seen_stems.update(d)
        stem_pools.append(by_t)
    targets = split_by_weight(sum(types), [u[2] for u in units])
    # 单元池总量不足：该单元按实有题数出卷，缺口并入 A 合集
    for c in range(1, len(units)):
        short = targets[c] - sum(len(d) for d in stem_pools[c])
        if short > 0:
            targets[c] -= short
            targets[0] += short
    if targets[0] > sum(len(d) for d in stem_pools[0]):
        raise ValueError(f"{pos['name']}: A 合集池不足（需 {targets[0]} 题，实有 {sum(len(d) for d in stem_pools[0])} 题）")
    caps = [[len(sp[t]) for sp in stem_pools] for t in range(len(types))]
    m = gen_matrix(types, targets, caps)
    paper, dist = [], []
    for c, (label, _, _) in enumerate(units):
        pick = []
        for t in range(len(types)):
            for stem in rng.sample(list(stem_pools[c][t]), m[t][c]):
                pick.append(rng.choice(stem_pools[c][t][stem]))
        paper.extend(pick)
        dist.append((label, len(pick)))
    paper.sort(key=lambda i: rows[i][1]['t'])  # 稳定排序：单选→多选→判断，单元顺序保留
    return paper, dist


def check_paper(name, rows, paper, types):
    """出卷闸门：题量、题型分布、无重复、答案字母不超选项范围。"""
    got = [0, 0, 0]
    seen = set()
    for i in paper:
        q = rows[i][1]
        got[q['t']] += 1
        if q['s'] in seen:
            raise ValueError(f'{name}: 题干重复 {q["s"][:40]}')
        seen.add(q['s'])
        if q['t'] == 2:
            if q['a'] not in ('对', '错'):
                raise ValueError(f'{name}: 判断题答案异常 {q["a"]!r}')
        else:
            bad = set(q['a']) - set(LETTERS[:len(q['o'])])
            if bad:
                raise ValueError(f'{name}: 答案字母 {q["a"]} 超出选项范围（{len(q["o"])} 个选项）')
            if q['t'] == 0 and len(q['a']) != 1:
                raise ValueError(f'{name}: 单选题答案非单字母 {q["a"]!r}')
            if q['t'] == 1 and len(q['a']) < 2:
                raise ValueError(f'{name}: 多选题答案不足 2 个字母 {q["a"]!r}')
    if got != list(types):
        raise ValueError(f'{name}: 题型题量 {got} != {list(types)}')


def write_paper(path, name, rows, paper):
    wb = Workbook()
    ws = wb.active
    ws.title = name[:31]
    ws.append(HEADERS)
    for i in paper:
        q = rows[i][1]
        if q['t'] == 2:
            opts, ans = ['A、正确', 'B、错误'], 'A' if q['a'] == '对' else 'B'
        else:
            opts = [f'{LETTERS[j]}、{o}' for j, o in enumerate(q['o'])]
            ans = q['a']
        ws.append([TYPE_NAME[q['t']], q['s'], *opts, *[''] * (8 - len(opts)), ans, q['n'], 1])
    for c in range(1, len(HEADERS) + 1):
        ws.cell(1, c).font = Font(bold=True)
    widths = {1: 10, 2: 60, 11: 10, 12: 50, 13: 6}
    for c in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 28)
    for r in range(2, ws.max_row + 1):
        for c in (2, 12):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description='按考试方案比例生成岗位试卷导入表')
    ap.add_argument('--positions', nargs='+', help='岗位名（默认 config.xlsx 全部岗位）')
    ap.add_argument('--types', nargs=3, type=int, default=[42, 43, 15],
                    metavar=('单选', '多选', '判断'), help='题型题量（默认 42 43 15）')
    ap.add_argument('--seed', type=int, help='随机种子（便于复现同一张卷）')
    args = ap.parse_args()

    rows = load_rows()
    positions = load_positions()
    if args.positions:
        by_name = {p['name']: p for p in positions}
        missing = [n for n in args.positions if n not in by_name]
        if missing:
            raise SystemExit(f'config.xlsx 无这些岗位: {missing}（现有: {[p["name"] for p in positions]}）')
        positions = [by_name[n] for n in args.positions]

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    rng = random.Random(args.seed)
    for pos in positions:
        paper, dist = gen_paper(pos, rows, args.types, rng)
        check_paper(pos['name'], rows, paper, args.types)
        path = OUT_DIR / f"{pos['name']}试卷.xlsx"
        write_paper(path, pos['name'], rows, paper)
        detail = ' '.join(f'{lb}={n}' for lb, n in dist)
        print(f"{pos['name']}: {path.name} 共 {len(paper)} 题（单选{args.types[0]}/多选{args.types[1]}/判断{args.types[2]}）")
        print(f"  单元分布: {detail}")


if __name__ == '__main__':
    main()
